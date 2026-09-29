"""
Lean Six Sigma process-observation timer for the food bank.

WHY this file is structured the way it is:
- All task names live in one place at the top (FLOW_TASKS / OTHER_TASKS) because
  the observer (not a programmer) will likely add/remove tasks after watching
  a real day and needs one obvious spot to edit.
- The Excel sheet is split into a "structure" (written ONCE, the first time the
  file is created) and "formulas" (REWRITTEN every save). This split exists so
  that hand-entered values -- the USL/LSL spec-limit cells the user fills in by
  hand in Excel -- never get clobbered when the app appends new rows and has to
  widen the formula ranges to match.
"""

import os
from datetime import datetime

import tkinter as tk
from tkinter import messagebox

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill, Font, Alignment

# ---------------------------------------------------------------------------
# EDIT HERE: the task list. Buttons are generated from these two lists, in
# this order, so re-ordering a list re-orders the buttons/summary rows too.
# ---------------------------------------------------------------------------

FLOW_TASKS = [
    "Walk: Cart<->Client",
    "Ask Client",
    "Walk: Client->Cart",
    "Tallying",
    "Grab Item",
    "Give Item",
    "Walk: Client->Client",
]

# The 3 flow steps that are physically walking -- used for the "total
# transportation time" roll-up. Kept as an explicit list (not auto-detected
# from the name) so it stays correct even if a walk step gets renamed.
WALK_TASKS = [
    "Walk: Cart<->Client",
    "Walk: Client->Cart",
    "Walk: Client->Client",
]

OTHER_TASKS = [
    "Take Donations",
    "Reload Supply",
    "Wrong Item (Redo)",
    "Distracted/Forgot Item",
    "Give Advice/Directions",
    "Help with Pet Food",
    "Help New Person Paperwork",
    "Egg Carton Return",
    "Chill Time/Break",
    "Get Rid of Boxes",
]

ALL_PRESET_TASKS = FLOW_TASKS + OTHER_TASKS

EXCEL_FILENAME = "mom_timer_log.xlsx"
SHEET_NAME = "Process Log"

INSTRUCTIONS_TEXT = (
    "Tap a task to start it. Pause stops the clock. Resume restarts the last task."
)

# ---------------------------------------------------------------------------
# Colors. Kept as named constants so the "what color means what" logic only
# has to be read once, in get_category() / fill_for_category() below.
# ---------------------------------------------------------------------------

COLOR_FLOW = "C6E0B4"     # light green  - one of the 7 core flow steps
COLOR_OTHER = "D9D9D9"    # light gray   - one of the 10 preset "other" tasks
COLOR_CUSTOM = "FFF2CC"   # light yellow - typed-in, not in either preset list
COLOR_INPUT = "FCE4D6"    # salmon       - "you need to type a number here"

FILL_FLOW = PatternFill(start_color=COLOR_FLOW, end_color=COLOR_FLOW, fill_type="solid")
FILL_OTHER = PatternFill(start_color=COLOR_OTHER, end_color=COLOR_OTHER, fill_type="solid")
FILL_CUSTOM = PatternFill(start_color=COLOR_CUSTOM, end_color=COLOR_CUSTOM, fill_type="solid")
FILL_INPUT = PatternFill(start_color=COLOR_INPUT, end_color=COLOR_INPUT, fill_type="solid")

BOLD = Font(bold=True)

# ---------------------------------------------------------------------------
# GUI button colors -- separate from the Excel fill colors above. The 7 flow
# steps are tapped constantly (the core repeated loop), so each gets its own
# distinct color to tap by shape/position/color without reading the label.
# The 10 "other" tasks are interrupts, tapped rarely, so they stay one plain
# gray -- a validated 10-way distinct-hue set doesn't exist, and consistent
# gray still reads instantly as "not part of the main loop."
# (Hues are a validated colorblind-safe categorical set -- 7 of 8 slots,
# skipping red since it's not being used to mean "stop/error" here.)
# ---------------------------------------------------------------------------

FLOW_BUTTON_COLORS = {
    "Walk: Cart<->Client": ("#2a78d6", "white"),   # blue
    "Ask Client": ("#eb6834", "black"),            # orange
    "Walk: Client->Cart": ("#1baf7a", "black"),    # aqua
    "Tallying": ("#eda100", "black"),               # yellow
    "Grab Item": ("#e87ba4", "black"),             # magenta
    "Give Item": ("#008300", "white"),             # green
    "Walk: Client->Client": ("#4a3aa7", "white"),  # violet
}
OTHER_BUTTON_COLOR = ("#d9d9d9", "black")  # matches the Excel "other" gray

# ---------------------------------------------------------------------------
# Sheet layout constants. Defined once, as math off the task-list lengths,
# so adding/removing a preset task automatically shifts everything below it
# instead of silently leaving stale rows behind.
# ---------------------------------------------------------------------------

RAW_DATA_START_ROW = 2           # row 1 is the raw-data header

SUMMARY_DATA_START_ROW = 2       # row 1 is the summary header
SUMMARY_DATA_END_ROW = SUMMARY_DATA_START_ROW + len(ALL_PRESET_TASKS) - 1
FLOW_SUMMARY_FIRST_ROW = SUMMARY_DATA_START_ROW
FLOW_SUMMARY_LAST_ROW = SUMMARY_DATA_START_ROW + len(FLOW_TASKS) - 1

LEGEND_START_ROW = SUMMARY_DATA_END_ROW + 3
DAY_SPLIT_START_ROW = LEGEND_START_ROW + 5

HELPER_HEADER_ROW = 1
HELPER_DATA_START_ROW = 2

CPK_BLOCK_1_START_ROW = 2    # "Cycle Time, CORE FLOW ONLY"       -> helper col P
CPK_BLOCK_2_START_ROW = 14   # "Cycle Time, INCLUDING INTERRUPTIONS" -> helper col Q


# ---------------------------------------------------------------------------
# Small helpers shared by GUI + Excel code
# ---------------------------------------------------------------------------

def get_category(task_name):
    """Which color bucket a task name falls into. Custom = not in either preset list."""
    if task_name in FLOW_TASKS:
        return "flow"
    if task_name in OTHER_TASKS:
        return "other"
    return "custom"


def fill_for_category(category):
    return {"flow": FILL_FLOW, "other": FILL_OTHER, "custom": FILL_CUSTOM}[category]


def format_elapsed(seconds, force_hours=False):
    """mm:ss for short tasks, h:mm:ss once we pass an hour (or always, for the
    session-total readout, which force_hours=True keeps on h:mm:ss from 0)."""
    seconds = int(seconds)
    hours, rem = divmod(seconds, 3600)
    minutes, secs = divmod(rem, 60)
    if hours or force_hours:
        return f"{hours}:{minutes:02d}:{secs:02d}"
    return f"{minutes:02d}:{secs:02d}"


# ---------------------------------------------------------------------------
# Excel: one-time structure (labels, headers, legend, static input cells)
# ---------------------------------------------------------------------------

def build_structure(ws):
    """Write everything that should NEVER be overwritten again, including the
    blank USL/LSL input cells the user fills in by hand in Excel."""

    # Raw data header (columns A-F)
    ws.append(["Cycle", "Task", "Note", "Start Time", "End Time", "Duration (sec)"])
    for col in range(1, 7):
        ws.cell(1, col).font = BOLD

    # Summary header (columns H-M)
    summary_headers = ["Task", "Count", "Total Time (sec)", "Total Time (min)",
                        "Avg Time (sec)", "% of Total Time"]
    for i, text in enumerate(summary_headers):
        cell = ws.cell(1, 8 + i, text)
        cell.font = BOLD

    # One summary row per preset task, pre-populated whether or not it has fired yet.
    for i, task in enumerate(ALL_PRESET_TASKS):
        row = SUMMARY_DATA_START_ROW + i
        ws.cell(row, 8, task)
        fill = FILL_FLOW if task in FLOW_TASKS else FILL_OTHER
        for col in range(8, 14):  # H..M
            ws.cell(row, col).fill = fill

    # Legend
    r = LEGEND_START_ROW
    ws.cell(r, 8, "Legend:").font = BOLD
    legend_rows = [
        (FILL_FLOW, "Core Flow Step (one of the 7 repeated-loop tasks)"),
        (FILL_OTHER, "Other / Preset Task"),
        (FILL_CUSTOM, "Custom / Unclassified Task (typed in, needs review)"),
    ]
    for offset, (fill, label) in enumerate(legend_rows, start=1):
        ws.cell(r + offset, 8).fill = fill
        ws.cell(r + offset, 9, label)

    # Day-level split labels (values are formulas, written in refresh_formulas)
    r = DAY_SPLIT_START_ROW
    ws.cell(r, 8, ("Day-Level Summary (excludes paused/idle time -- "
                   "pauses are never logged as rows)")).font = BOLD
    ws.cell(r + 1, 8, "Total Logged Time (sec)")
    ws.cell(r + 2, 8, "Total Core Flow Time (sec)")
    ws.cell(r + 3, 8, "Total \"Extras\" Time (sec)")
    ws.cell(r + 4, 8, "Total Core Flow Time (% of day)")
    ws.cell(r + 5, 8, "Total \"Extras\" Time (% of day)")
    ws.cell(r + 6, 8, "Total Transportation Time -- All Walking (sec)")
    ws.cell(r + 7, 8, "Total Transportation Time -- All Walking (% of day)")

    # Helper-table header (columns O:Q)
    ws.cell(HELPER_HEADER_ROW, 15, "Cycle").font = BOLD
    ws.cell(HELPER_HEADER_ROW, 16, "Core Flow Time (sec)").font = BOLD
    ws.cell(HELPER_HEADER_ROW, 17, "Total Time (sec)").font = BOLD

    # Cpk blocks (labels + blank, shaded spec-limit input cells only)
    write_cpk_block_labels(ws, CPK_BLOCK_1_START_ROW, "Cycle Time, CORE FLOW ONLY")
    write_cpk_block_labels(ws, CPK_BLOCK_2_START_ROW, "Cycle Time, INCLUDING INTERRUPTIONS")

    # Cosmetic column widths -- purely so it's readable without manual resizing.
    widths = {"A": 8, "B": 22, "C": 22, "D": 19, "E": 19, "F": 14,
              "H": 28, "I": 36, "S": 34, "T": 16}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width


def write_cpk_block_labels(ws, start_row, title):
    """The reusable Cpk block 'shape' -- labels only. Called twice (once per
    block) so both blocks are guaranteed to stay structurally identical."""
    ws.merge_cells(f"S{start_row}:T{start_row}")
    title_cell = ws.cell(start_row, 19, title)
    title_cell.font = BOLD
    title_cell.alignment = Alignment(horizontal="center")

    labels = [
        "Upper Spec Limit (USL, sec)",
        "Lower Spec Limit (LSL, sec) [optional]",
        "Cycle Count",
        "Mean Cycle Time (sec)",
        "Std Dev Cycle Time (sec)",
        "Cpu",
        "Cpl",
        "Cpk",
        "Estimated Sigma Level (approx, = 3 x Cpk)",
    ]
    for offset, label in enumerate(labels, start=1):
        ws.cell(start_row + offset, 19, label)

    # USL/LSL are the only hand-entered cells on this whole sheet -- shade them
    # so it's obvious input is needed, and leave the value blank on purpose.
    ws.cell(start_row + 1, 20).fill = FILL_INPUT
    ws.cell(start_row + 2, 20).fill = FILL_INPUT


# ---------------------------------------------------------------------------
# Excel: formulas that depend on how many rows/cycles exist so far.
# Rewritten every save -- safe, because nothing hand-entered lives in these
# cells (the only hand-entered cells, USL/LSL, are skipped here on purpose).
# ---------------------------------------------------------------------------

def refresh_formulas(ws, last_data_row, cycle_count):
    lr = max(last_data_row, RAW_DATA_START_ROW)  # keep ranges valid even with 0 rows logged
    rng_task = f"$B$2:$B${lr}"
    rng_dur = f"$F$2:$F${lr}"
    rng_cycle = f"$A$2:$A${lr}"

    _refresh_summary_block(ws, rng_task, rng_dur)
    _refresh_day_split(ws, rng_task, rng_dur)
    last_helper_row = _refresh_helper_table(ws, cycle_count, rng_task, rng_dur, rng_cycle)
    refresh_cpk_block(ws, CPK_BLOCK_1_START_ROW, "P", last_helper_row)
    refresh_cpk_block(ws, CPK_BLOCK_2_START_ROW, "Q", last_helper_row)


def _refresh_summary_block(ws, rng_task, rng_dur):
    for i in range(len(ALL_PRESET_TASKS)):
        row = SUMMARY_DATA_START_ROW + i
        task_cell = f"$H${row}"
        ws.cell(row, 9, f"=COUNTIF({rng_task},{task_cell})")                       # Count
        ws.cell(row, 10, f"=SUMIF({rng_task},{task_cell},{rng_dur})")              # Total sec
        ws.cell(row, 11, f"=J{row}/60")                                            # Total min
        ws.cell(row, 12, f"=IF(I{row}=0,0,J{row}/I{row})")                         # Avg sec
        pct_cell = ws.cell(row, 13, f'=IF(SUM({rng_dur})=0,0,J{row}/SUM({rng_dur}))')
        pct_cell.number_format = "0.0%"


def _refresh_day_split(ws, rng_task, rng_dur):
    r = DAY_SPLIT_START_ROW
    ws.cell(r + 1, 9, f"=SUM({rng_dur})")
    ws.cell(r + 2, 9, f"=SUM($J${FLOW_SUMMARY_FIRST_ROW}:$J${FLOW_SUMMARY_LAST_ROW})")
    ws.cell(r + 3, 9, f"=I{r + 1}-I{r + 2}")
    ws.cell(r + 4, 9, f"=IF(I{r + 1}=0,0,I{r + 2}/I{r + 1})").number_format = "0.0%"
    ws.cell(r + 5, 9, f"=IF(I{r + 1}=0,0,I{r + 3}/I{r + 1})").number_format = "0.0%"

    # Total transportation time: sum of just the 3 walking tasks, wherever
    # they land in the raw data (not a contiguous summary-row range, since
    # the other 4 flow steps sit between them).
    walk_terms = " + ".join(f'SUMIF({rng_task},"{task}",{rng_dur})' for task in WALK_TASKS)
    ws.cell(r + 6, 9, "=" + walk_terms)
    ws.cell(r + 7, 9, f"=IF(I{r + 1}=0,0,I{r + 6}/I{r + 1})").number_format = "0.0%"


def _refresh_helper_table(ws, cycle_count, rng_task, rng_dur, rng_cycle):
    """Columns O:Q -- one row per cycle number, flow-only time and all-task time.
    Cleared a few rows past the current max so a rare shrink never leaves
    stale numbers behind."""
    n = max(cycle_count, 1)
    clear_past = n + 3
    for c in range(1, clear_past + 1):
        row = HELPER_DATA_START_ROW + c - 1
        if c > n:
            ws.cell(row, 15).value = None
            ws.cell(row, 16).value = None
            ws.cell(row, 17).value = None
            continue
        ws.cell(row, 15, c)
        flow_terms = " + ".join(
            f'SUMIFS({rng_dur},{rng_cycle},$O${row},{rng_task},"{task}")'
            for task in FLOW_TASKS
        )
        ws.cell(row, 16, "=" + flow_terms)
        ws.cell(row, 17, f"=SUMIF({rng_cycle},$O${row},{rng_dur})")
    return HELPER_DATA_START_ROW + n - 1


def refresh_cpk_block(ws, start_row, helper_col_letter, last_helper_row):
    """The reusable Cpk formula pattern -- same shape for both blocks, just
    pointed at a different helper column (P = core flow, Q = incl. interruptions)."""
    usl = f"T{start_row + 1}"
    lsl = f"T{start_row + 2}"
    count_c = f"T{start_row + 3}"
    mean_c = f"T{start_row + 4}"
    std_c = f"T{start_row + 5}"
    cpu_c = f"T{start_row + 6}"
    cpl_c = f"T{start_row + 7}"
    cpk_c = f"T{start_row + 8}"
    sigma_c = f"T{start_row + 9}"
    data_rng = f"{helper_col_letter}2:{helper_col_letter}{last_helper_row}"

    ws[count_c] = f"=COUNT({data_rng})"
    ws[mean_c] = f'=IF(COUNT({data_rng})=0,"",AVERAGE({data_rng}))'
    # STDEV (not STDEV.S) on purpose -- identical sample stdev, but STDEV.S can
    # throw #NAME? when written by non-Excel tools without an _xlfn. prefix.
    ws[std_c] = f'=IF(COUNT({data_rng})<2,"",STDEV({data_rng}))'

    # ISBLANK, not IFERROR: a blank spec-limit cell evaluates as 0 in a formula,
    # which would silently produce a wrong (not missing) Cpu/Cpl.
    ws[cpu_c] = (
        f'=IF(ISBLANK({usl}),"",IF(OR({mean_c}="",{std_c}="",{std_c}=0),"",'
        f"({usl}-{mean_c})/(3*{std_c})))"
    )
    ws[cpl_c] = (
        f'=IF(ISBLANK({lsl}),"",IF(OR({mean_c}="",{std_c}="",{std_c}=0),"",'
        f"({mean_c}-{lsl})/(3*{std_c})))"
    )
    ws[cpk_c] = (
        f'=IF(AND({cpu_c}="",{cpl_c}=""),"Enter at least one spec limit",'
        f'IF({cpu_c}="",{cpl_c},IF({cpl_c}="",{cpu_c},MIN({cpu_c},{cpl_c}))))'
    )
    ws[sigma_c] = f'=IF(ISNUMBER({cpk_c}),3*{cpk_c},"")'


# ---------------------------------------------------------------------------
# Excel: appending a completed task + loading/creating the workbook
# ---------------------------------------------------------------------------

def append_raw_row(ws, row, cycle, task, note, start_dt, end_dt, duration):
    ws.cell(row, 1, cycle)
    ws.cell(row, 2, task)
    ws.cell(row, 3, note or None)
    start_cell = ws.cell(row, 4, start_dt)
    start_cell.number_format = "yyyy-mm-dd hh:mm:ss"
    end_cell = ws.cell(row, 5, end_dt)
    end_cell.number_format = "yyyy-mm-dd hh:mm:ss"
    ws.cell(row, 6, round(duration, 2))

    fill = fill_for_category(get_category(task))
    for col in range(1, 7):
        ws.cell(row, col).fill = fill


def find_last_data_row(ws):
    """Scan column A instead of trusting ws.max_row, in case of stray blank rows."""
    last_row = 1
    for row in range(RAW_DATA_START_ROW, ws.max_row + 1):
        if ws.cell(row, 1).value is not None:
            last_row = row
    return last_row


def find_max_cycle(ws, last_data_row):
    max_cycle = 0
    for row in range(RAW_DATA_START_ROW, last_data_row + 1):
        value = ws.cell(row, 1).value
        if isinstance(value, (int, float)) and value > max_cycle:
            max_cycle = value
    return int(max_cycle)


def ensure_workbook(path):
    """Load the existing file (appending to past sessions) or create a new one."""
    if os.path.exists(path):
        wb = load_workbook(path)
        if SHEET_NAME in wb.sheetnames:
            ws = wb[SHEET_NAME]
            return wb, ws, False
        ws = wb.create_sheet(SHEET_NAME)
        build_structure(ws)
        return wb, ws, True

    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    build_structure(ws)
    return wb, ws, True


# ---------------------------------------------------------------------------
# The app itself
# ---------------------------------------------------------------------------

class TimerApp:
    def __init__(self, root):
        self.root = root
        self.root.title("Food Bank Process Timer")

        script_dir = os.path.dirname(os.path.abspath(__file__))
        self.excel_path = os.path.join(script_dir, EXCEL_FILENAME)

        self.wb, self.ws, created = ensure_workbook(self.excel_path)
        self.last_data_row = find_last_data_row(self.ws)
        self.cycle_number = max(find_max_cycle(self.ws, self.last_data_row), 1)

        self.current_task = None
        self.current_task_start = None
        self.last_paused_task = None

        # "Total time elapsed" readout: accumulated seconds from segments that
        # were already closed out, plus whatever's running right now. It only
        # starts once the first task is clicked, and pauses along with Pause.
        self.total_accum = 0.0
        self.total_segment_start = None

        # Write the structure + first set of formulas to disk right away so the
        # file exists and is inspectable even before the first task is logged.
        self._save_and_refresh()

        self._build_gui()
        self.root.protocol("WM_DELETE_WINDOW", self._on_close)
        self._tick()

    # -- core timer logic ---------------------------------------------------

    def _stop_and_log_current(self):
        """Close out whatever's running: write its row, clear the Note field.
        Returns True if something was actually logged."""
        if self.current_task is None:
            return False

        end = datetime.now()
        duration = (end - self.current_task_start).total_seconds()
        note = self.note_var.get().strip()

        self.last_data_row += 1
        append_raw_row(self.ws, self.last_data_row, self.cycle_number,
                       self.current_task, note, self.current_task_start, end, duration)

        self.note_var.set("")
        self.current_task = None
        self.current_task_start = None
        return True

    def _start_task(self, name):
        """Shared by preset buttons, custom task, and Resume: stop whatever's
        running (if anything), then start the new one. Same button twice = no-op."""
        if self.current_task == name:
            return
        logged = self._stop_and_log_current()
        self.current_task = name
        self.current_task_start = datetime.now()
        if self.total_segment_start is None:  # first task, or resuming after a pause
            self.total_segment_start = self.current_task_start
        self._update_running_label()
        if logged:
            self._save_and_refresh()

    def _pause(self):
        if self.current_task is None:
            return
        name = self.current_task
        if self._stop_and_log_current():
            self.last_paused_task = name
            self.resume_button.config(state="normal")
            self._save_and_refresh()
        if self.total_segment_start is not None:
            self.total_accum += (datetime.now() - self.total_segment_start).total_seconds()
            self.total_segment_start = None
        self._update_running_label()

    def _resume(self):
        if self.last_paused_task is None:
            return
        self._start_task(self.last_paused_task)

    def _new_cycle(self):
        self.cycle_number += 1
        self.cycle_var.set(f"Cycle: {self.cycle_number}")
        self._save_and_refresh()

    def _start_custom_task(self):
        name = self.custom_var.get().strip()
        if not name:
            return
        self.custom_var.set("")
        self._start_task(name)

    # -- Excel persistence ---------------------------------------------------

    def _save_and_refresh(self):
        refresh_formulas(self.ws, self.last_data_row, self.cycle_number)
        try:
            self.wb.save(self.excel_path)
        except PermissionError:
            messagebox.showerror(
                "File In Use",
                f"Could not save {EXCEL_FILENAME} -- it looks like the file is "
                "currently open in Excel.\n\nPlease close it there. Nothing was "
                "lost: this entry is still in memory and will be saved on the "
                "next action.",
            )

    # -- GUI -------------------------------------------------------------

    def _build_gui(self):
        big_font = ("Arial", 14)
        button_font = ("Arial", 11)

        status_frame = tk.Frame(self.root, pady=8)
        status_frame.pack(fill="x")
        self.session_elapsed_var = tk.StringVar(value="00:00:00")
        self.running_var = tk.StringVar(value="Idle")
        self.elapsed_var = tk.StringVar(value="00:00")
        self.cycle_var = tk.StringVar(value=f"Cycle: {self.cycle_number}")
        tk.Label(status_frame, text="Total Time Elapsed (this session):", font=("Arial", 9), fg="gray30").pack()
        tk.Label(status_frame, textvariable=self.session_elapsed_var, font=("Arial", 11)).pack()
        tk.Label(status_frame, textvariable=self.running_var, font=big_font).pack(pady=(8, 0))
        tk.Label(status_frame, textvariable=self.elapsed_var, font=("Arial", 22, "bold")).pack()
        tk.Label(status_frame, textvariable=self.cycle_var, font=big_font).pack()

        entry_frame = tk.Frame(self.root, pady=6)
        entry_frame.pack(fill="x", padx=10)

        tk.Label(entry_frame, text="Note:", font=button_font).grid(row=0, column=0, sticky="w")
        self.note_var = tk.StringVar()
        tk.Entry(entry_frame, textvariable=self.note_var, font=button_font, width=30).grid(
            row=0, column=1, sticky="we", padx=5)

        tk.Label(entry_frame, text="Custom Task:", font=button_font).grid(row=1, column=0, sticky="w")
        self.custom_var = tk.StringVar()
        tk.Entry(entry_frame, textvariable=self.custom_var, font=button_font, width=30).grid(
            row=1, column=1, sticky="we", padx=5)
        tk.Button(entry_frame, text="Start Custom Task", font=button_font,
                  command=self._start_custom_task).grid(row=1, column=2, padx=5)
        entry_frame.columnconfigure(1, weight=1)

        groups_frame = tk.Frame(self.root)
        groups_frame.pack(fill="both", expand=True, padx=10, pady=6)

        flow_frame = tk.LabelFrame(groups_frame, text="Repeated Flow Steps", font=button_font, padx=6, pady=6)
        flow_frame.pack(side="left", fill="both", expand=True, padx=(0, 5))
        self._add_task_buttons(flow_frame, FLOW_TASKS, columns=1, font=button_font,
                                color_lookup=FLOW_BUTTON_COLORS)

        other_frame = tk.LabelFrame(groups_frame, text="Other Tasks", font=button_font, padx=6, pady=6)
        other_frame.pack(side="left", fill="both", expand=True, padx=(5, 0))
        self._add_task_buttons(other_frame, OTHER_TASKS, columns=2, font=button_font,
                                color_lookup=None)

        control_frame = tk.Frame(self.root, pady=8)
        control_frame.pack(fill="x", padx=10)
        tk.Button(control_frame, text="Pause", font=button_font, width=12,
                  command=self._pause).pack(side="left", padx=5)
        self.resume_button = tk.Button(control_frame, text="Resume", font=button_font, width=12,
                                        command=self._resume, state="disabled")
        self.resume_button.pack(side="left", padx=5)
        tk.Button(control_frame, text="New Cycle", font=button_font, width=12,
                  command=self._new_cycle).pack(side="left", padx=5)

        tk.Label(self.root, text=INSTRUCTIONS_TEXT, font=("Arial", 10), fg="gray20").pack(
            side="bottom", pady=6)

    def _add_task_buttons(self, parent, task_list, columns, font, color_lookup):
        for i, task in enumerate(task_list):
            row, col = divmod(i, columns)
            bg, fg = color_lookup.get(task, OTHER_BUTTON_COLOR) if color_lookup else OTHER_BUTTON_COLOR
            btn = tk.Button(parent, text=task, font=font, width=20, height=2,
                             wraplength=150, bg=bg, fg=fg, activebackground=bg,
                             command=lambda t=task: self._start_task(t))
            btn.grid(row=row, column=col, padx=4, pady=4, sticky="we")
        for col in range(columns):
            parent.columnconfigure(col, weight=1)

    def _update_running_label(self):
        if self.current_task:
            self.running_var.set(f"Running: {self.current_task}")
        else:
            self.running_var.set("Idle (paused)" if self.last_paused_task else "Idle")

    def _tick(self):
        total_elapsed = self.total_accum
        if self.total_segment_start is not None:
            total_elapsed += (datetime.now() - self.total_segment_start).total_seconds()
        self.session_elapsed_var.set(format_elapsed(total_elapsed, force_hours=True))

        if self.current_task is not None:
            elapsed = (datetime.now() - self.current_task_start).total_seconds()
            self.elapsed_var.set(format_elapsed(elapsed))
        else:
            self.elapsed_var.set("00:00")
        self.root.after(500, self._tick)

    def _on_close(self):
        if self.current_task is not None:
            if messagebox.askyesno(
                "Task Still Running",
                f'"{self.current_task}" is still running. Stop and log it before closing?',
            ):
                self._stop_and_log_current()
                self._save_and_refresh()
        self.root.destroy()


if __name__ == "__main__":
    root = tk.Tk()
    app = TimerApp(root)
    root.mainloop()
